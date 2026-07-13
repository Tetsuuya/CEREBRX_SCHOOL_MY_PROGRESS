import openpyxl

src_file = "Attendance_Sample.xlsx"
dest_file = "Clean_Attendace_template_v3.xlsx"

print(f"Loading sample file: {src_file}...")
wb = openpyxl.load_workbook(src_file)
sheet = wb.active

# 1. Clear columns in Row 8 (Male template row)
sheet['A8'] = None  # Clear index
sheet['B8'] = None  # Clear name
sheet['C8'] = None  # Clear Flag
sheet['D8'] = None  # Clear Chapel
sheet['E8'] = None  # Clear manual columns
sheet['F8'] = None
sheet['G8'] = None
sheet['H8'] = None  # Clear Prayer
sheet['I8'] = None  # Clear Amount

# Unmerge A29:B29 before deleting rows to prevent dangling merged cells in openpyxl
print("Unmerging A29:B29...")
sheet.unmerge_cells('A29:B29')

# 2. Delete rows 9 to 28 (old male students sample data).
# Since 20 rows are deleted starting at row 9, the old row 29 ("Female" header) moves to row 9,
# and the old row 30 (first female student) moves to row 10.
print("Deleting old male student rows (9 to 28)...")
sheet.delete_rows(idx=9, amount=20)

# Merge A9:B9 for the new Female header position
print("Merging A9:B9...")
sheet.merge_cells('A9:B9')

# Clear columns C to I in the "Female" header row (now Row 9) to make it consistent with Row 7 (Male)
print("Cleaning columns C to I in Row 9 (Female header)...")
for col_idx in range(3, 10):
    cell = sheet.cell(row=9, column=col_idx)
    cell.value = None
    cell.border = openpyxl.styles.Border()
    cell.fill = openpyxl.styles.PatternFill(fill_type=None)

# 3. Clear columns in Row 10 (Female template row)
sheet['A10'] = None  # Clear index
sheet['B10'] = None  # Clear name
sheet['C10'] = None  # Clear Flag
sheet['D10'] = None  # Clear Chapel
sheet['E10'] = None  # Clear manual columns
sheet['F10'] = None
sheet['G10'] = None
sheet['H10'] = None  # Clear Prayer
sheet['I10'] = None  # Clear Amount

# 4. Delete all rows from 11 onwards (old female student sample data)
print("Deleting old female student rows (11 onwards)...")
sheet.delete_rows(idx=11, amount=100)

# 5. Delete Columns J to Q (Columns 10 to 17) to remove the right side completely
print("Deleting columns J to Q...")
sheet.delete_cols(idx=10, amount=8)

# Save as the clean template
print(f"Saving cleaned template as: {dest_file}...")
wb.save(dest_file)
print("Clean template version 3 (no right side) successfully generated!")
