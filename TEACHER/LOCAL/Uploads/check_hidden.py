import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Rhenel Jhon Sajol\Desktop\CEREB_SCHOOL_BACKUP\TEACHER\LOCAL\Uploads\3-term_New_Grade_Template_ULTRA_OPTIMIZED.xlsx')

print("Checking hidden rows in all sheets:")
print("=" * 60)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    hidden = [row for row in range(1, 51) if row in ws.row_dimensions and ws.row_dimensions[row].hidden]
    
    if hidden:
        print(f"\nSheet: {sheet_name}")
        print(f"Hidden rows: {hidden}")
        
        # Check content of hidden rows
        for row_num in hidden:
            print(f"  Row {row_num} content:")
            row_cells = ws[row_num]
            has_content = False
            for cell in row_cells:
                if cell.value is not None and str(cell.value).strip():
                    print(f"    {cell.coordinate}: {repr(cell.value)}")
                    has_content = True
            if not has_content:
                print(f"    (EMPTY ROW)")

wb.close()
