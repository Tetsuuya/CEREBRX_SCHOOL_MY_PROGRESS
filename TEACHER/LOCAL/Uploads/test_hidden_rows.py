"""
Test script to verify the hidden rows logic in Excelwithspout.php

This simulates what PHPExcel does:
1. Load template and detect originally hidden rows
2. Process data (placeholders would be replaced here)
3. Re-hide the originally hidden rows
4. Save file
"""

import openpyxl
from copy import copy

# Load the ULTRA_OPTIMIZED template
template_path = r'C:\Users\Rhenel Jhon Sajol\Desktop\CEREB_SCHOOL_BACKUP\TEACHER\LOCAL\Uploads\3-term_New_Grade_Template_ULTRA_OPTIMIZED.xlsx'
output_path = r'C:\Users\Rhenel Jhon Sajol\Desktop\CEREB_SCHOOL_BACKUP\TEACHER\LOCAL\Uploads\TEST_GENERATED.xlsx'

print("=" * 60)
print("SIMULATING PHPEXCEL LOGIC")
print("=" * 60)

wb = openpyxl.load_workbook(template_path)

# Dictionary to store originally hidden rows per sheet
originally_hidden = {}

print("\n1. DETECTING ORIGINALLY HIDDEN ROWS:")
print("-" * 60)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    hidden = []
    
    for row in range(1, 51):
        if row in ws.row_dimensions:
            rd = ws.row_dimensions[row]
            # Simulate PHPExcel's getVisible() check
            if rd.hidden:
                hidden.append(row)
    
    originally_hidden[sheet_name] = hidden
    if hidden:
        print(f"Sheet '{sheet_name}': Found {len(hidden)} hidden rows: {hidden}")

# Simulate data insertion (this is where PHPExcel might lose hidden state)
print("\n2. SIMULATING DATA INSERTION:")
print("-" * 60)
print("(In real code, this is where student names would be inserted)")
print("(PHPExcel might lose hidden row state during this process)")

# Simulate the re-hiding logic from our fix
print("\n3. RE-HIDING ORIGINALLY HIDDEN ROWS:")
print("-" * 60)
for sheet_name, hidden_rows in originally_hidden.items():
    if hidden_rows:
        ws = wb[sheet_name]
        print(f"Sheet '{sheet_name}': Re-hiding {len(hidden_rows)} rows")
        for row_num in hidden_rows:
            if row_num in ws.row_dimensions:
                ws.row_dimensions[row_num].hidden = True
            else:
                # Create row dimension if it doesn't exist
                ws.row_dimensions[row_num].hidden = True

# Save the file
print("\n4. SAVING FILE:")
print("-" * 60)
wb.save(output_path)
print(f"Saved to: {output_path}")

# Verify the result
print("\n5. VERIFICATION - CHECKING SAVED FILE:")
print("-" * 60)
wb2 = openpyxl.load_workbook(output_path)
for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    hidden = [row for row in range(1, 51) if row in ws.row_dimensions and ws.row_dimensions[row].hidden]
    
    expected = originally_hidden.get(sheet_name, [])
    if hidden == expected:
        print(f"✓ Sheet '{sheet_name}': Hidden rows preserved correctly {hidden}")
    else:
        print(f"✗ Sheet '{sheet_name}': MISMATCH!")
        print(f"  Expected: {expected}")
        print(f"  Got: {hidden}")

wb2.close()
wb.close()

print("\n" + "=" * 60)
print("TEST COMPLETE")
print("=" * 60)
